from openpyxl import Workbook, load_workbook

from config import INPUT_DIR, OUTPUT_DIR
from main import process


def write_invoice(fact):
    output = OUTPUT_DIR / "Facture_results.xlsx"
    exist = output.exists()
    if exist:
        wb = load_workbook(output)
        ws = wb.active
        ws.append(
            [
                fact["numero"],
                fact["fecha"],
                fact["emisor"],
                fact["total"],
                fact["formato"],
            ]
        )
        wb.save(output)
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Numero", "Fecha", "Emisor", "Total", "Formato"])
        ws.append(
            [
                fact["numero"],
                fact["fecha"],
                fact["emisor"],
                fact["total"],
                fact["formato"],
            ]
        )
        wb.save(output)
    return True


if __name__ == "__main__":
    pdf_path = INPUT_DIR / "factura_es_01.pdf"
    facture = process(pdf_path)
    write_invoice(facture)
